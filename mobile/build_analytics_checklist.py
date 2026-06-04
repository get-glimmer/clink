#!/usr/bin/env python3
"""Generate the analytics verification checklist as an .xlsx workbook."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()

# ---- styling helpers ----
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, size=11, color="1F4E78")
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="top", horizontal="center")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

COLUMNS = ["Action / Button", "Event(s) fired", "Screen", "Where to verify",
           "Sent to", "Status", "Notes"]

# rows: tuple(section_title) OR list[action, events, screen, verify, sent_to, notes]
ROWS = [
    ("1. App launch (pre-login)",),
    ["Cold-launch the app", "app_opened, session_started", "—",
     "Firebase DebugView", "Firebase only", "No user_events row yet (no session)"],

    ("2. Sign up / Sign in",),
    ["Tap 'Sign in with Apple'", "sign_up_started (source: apple)", "SignUp",
     "Firebase DebugView", "Firebase only", "Fires before the OS dialog opens"],
    ["Tap 'Sign in with Google'", "sign_up_started (source: google)", "SignUp",
     "Firebase DebugView", "Firebase only", "Fires before the OS dialog opens"],
    ["Email signup -> 'Continue' on Finish-signup", "sign_up_completed (source: email)", "SignUp",
     "Firebase DebugView", "Firebase only", ""],
    ["Land on Finish-signup screen", "screen_viewed, onboarding_started", "Onboarding",
     "Firebase DebugView", "Firebase only", "Onboarding only for OAuth profile completion"],
    ["Complete profile (OAuth) -> 'Continue'", "onboarding_completed", "Onboarding",
     "Firebase + Supabase", "Both", "Session now exists -> check user_events too"],
    ["Finish login (email / Apple / Google)", "sign_in_completed (source matches)", "—",
     "Firebase + Supabase", "Both", ""],
    ["Permission popup -> 'Allow' / 'Don't Allow'",
     "notification_permission_requested, then _granted or _denied", "—",
     "Firebase + Supabase", "Both", ""],

    ("3. Tasks & Categories (Home tab)",),
    ["Open Home tab", "screen_viewed", "Home", "Firebase + Supabase", "Both", ""],
    ["Options (...) -> 'Categories'", "category_opened, screen_viewed", "CategoryList",
     "Firebase + Supabase", "Both", ""],
    ["Add category -> 'Done'", "category_created (+ first_category_created)", "Home",
     "Firebase + Supabase", "Both", "first_* fires once per install"],
    ["Edit a category -> 'Done'", "category_edited", "Home", "Firebase + Supabase", "Both", ""],
    ["Delete a category (trash)", "category_deleted", "Home", "Firebase + Supabase", "Both", ""],
    ["Add a task -> save", "task_created (+ first_task_created)", "Home",
     "Firebase + Supabase", "Both", "first_* fires once per install"],
    ["Check a task's checkbox", "task_completed (+ starter_task_completed)", "Home",
     "Firebase + Supabase", "Both", "starter_* fires once per install"],
    ["Tap a task's '...' menu", "task_opened", "Home", "Firebase + Supabase", "Both", ""],
    ["Delete a task", "task_deleted", "Home", "Firebase + Supabase", "Both", ""],

    ("4. Friends / Community",),
    ["Open Community tab", "screen_viewed", "FriendList", "Firebase + Supabase", "Both", ""],
    ["Search a user -> send friend request", "friend_added", "FriendList",
     "Firebase + Supabase", "Both", ""],
    ["Accept an incoming request", "invite_accepted (source: friend_request)", "FriendList",
     "Firebase + Supabase", "Both", ""],
    ["Tap a friend to open their profile", "friend_profile_viewed", "FriendList",
     "Firebase + Supabase", "Both", ""],
    ["Send a text note to a friend", "friend_note_sent", "FriendList",
     "Firebase + Supabase", "Both", "Confirm NO note text in event_properties"],
    ["React with an emoji", "emoji_reacted", "FriendList",
     "Firebase + Supabase", "Both", "Confirm NO content in event_properties"],

    ("5. Challenges",),
    ["Open a challenge", "challenge_viewed, screen_viewed", "ChallengeDetail",
     "Firebase + Supabase", "Both", ""],
    ["Create challenge w/ friends invited -> submit", "challenge_created, invite_sent (per friend)",
     "ChallengeDetail", "Firebase + Supabase", "Both", ""],
    ["Accept a challenge invite -> 'Join'", "challenge_joined, invite_accepted", "ChallengeDetail",
     "Firebase + Supabase", "Both", ""],
    ["Check some challenge tasks -> submit day", "challenge_task_completed (with completed_count)",
     "ChallengeDetail", "Firebase + Supabase", "Both", ""],
    ["Check ALL of today's tasks -> submit", "challenge_completed", "ChallengeDetail",
     "Firebase + Supabase", "Both", "Once per challenge per day"],
    ["Leave a challenge", "challenge_left", "ChallengeDetail", "Firebase + Supabase", "Both", ""],

    ("6. Notifications",),
    ["Receive a push while app is open", "notification_received (with notification_type)", "—",
     "Firebase + Supabase", "Both", ""],
    ["Tap a push notification", "notification_clicked", "—", "Firebase + Supabase", "Both", ""],
]

# ================= Sheet 1: Checklist =================
ws = wb.active
ws.title = "Checklist"

# header row
for c, name in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    cell.border = BORDER

# status dropdown
dv = DataValidation(type="list", formula1='"Pass,Fail,N/A,Not tested"', allow_blank=True)
ws.add_data_validation(dv)

r = 2
for row in ROWS:
    if len(row) == 1:  # section header
        ws.cell(row=r, column=1, value=row[0])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLUMNS))
        cell = ws.cell(row=r, column=1)
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
        cell.alignment = Alignment(vertical="center")
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
        continue

    action, events, screen, verify, sent_to, notes = row
    values = [action, events, screen, verify, sent_to, "", notes]  # Status left blank
    for c, val in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.alignment = CENTER if c in (3, 5, 6) else WRAP
        cell.border = BORDER
    dv.add(ws.cell(row=r, column=6))  # Status dropdown
    r += 1

widths = [42, 40, 15, 20, 14, 13, 38]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False

# ================= Sheet 2: Setup =================
ws2 = wb.create_sheet("Setup & Notes")
setup = [
    ("Setup & verification notes", True),
    ("", False),
    ("Firebase DebugView", True),
    ("Xcode -> Edit Scheme -> Run -> Arguments -> add  -FIRDebugEnabled", False),
    ("Then: Firebase Console -> Analytics -> DebugView (events appear in seconds).", False),
    ("", False),
    ("Supabase", True),
    ("Supabase Dashboard -> Table Editor -> user_events. Refresh after each action.", False),
    ("", False),
    ("Routing rule (spec #8)", True),
    ("Pre-login events  -> Firebase ONLY (no session, no user_events row).", False),
    ("Post-login events -> Firebase + Supabase user_events.", False),
    ("", False),
    ("Columns to confirm on a post-login row", True),
    ("user_id, anonymous_id, session_id (same across one launch),", False),
    ("screen_name, previous_screen, source,", False),
    ("platform=ios, app_version, device_model (e.g. iPhone17,1), os_version,", False),
    ("event_properties = only ids/source -- NEVER names, notes, or messages (spec #7).", False),
    ("", False),
    ("once-per-install events", True),
    ("first_category_created, first_task_created, starter_task_completed fire ONCE per install.", False),
    ("To re-test, delete & reinstall the app.", False),
    ("challenge_completed dedupes once per challenge per day.", False),
    ("", False),
    ("Fastest smoke test (2 min)", True),
    ("1. Launch -> app_opened in Firebase DebugView.", False),
    ("2. Log in -> sign_in_completed in BOTH Firebase and Supabase.", False),
    ("3. Create a task -> task_created + first_task_created in both.", False),
    ("", False),
    ("If a row shows in Firebase but not Supabase -> Supabase insert issue (RLS / session).", False),
    ("If nothing shows in Firebase -> -FIRDebugEnabled flag or GoogleService-Info.plist target membership.", False),
]
for i, (text, is_head) in enumerate(setup, start=1):
    cell = ws2.cell(row=i, column=1, value=text)
    cell.font = Font(bold=True, size=12, color="1F4E78") if is_head else Font(size=11)
ws2.column_dimensions["A"].width = 100
ws2.sheet_view.showGridLines = False

out = "analytics_verification_checklist.xlsx"
wb.save(out)
print("wrote", out)
